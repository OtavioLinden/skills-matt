#!/usr/bin/env python3
"""Short summary of a received file; full text goes to a file, never stdout.

Usage: ~/.local/share/ler-documento/venv/bin/python -I resumo.py FILE [--out DIR] [--head N]
Never extracts archives and never executes anything from the input.
"""
import argparse
import csv
import hashlib
import shutil
import subprocess
import sys
from pathlib import Path

ARCH = {".zip", ".rar", ".7z", ".tar", ".gz", ".tgz", ".bz2", ".xz"}
IMG = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
MEDIA = {".mp3", ".wav", ".m4a", ".ogg", ".opus", ".flac", ".mp4", ".mkv", ".mov", ".webm", ".avi"}


def out_path(src: Path, out_dir: Path, ext: str) -> Path:
    tag = hashlib.sha1(str(src.resolve()).encode()).hexdigest()[:6]
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / f"{src.stem[:40]}-{tag}{ext}"


def clip(s, n=110):
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[: n - 1] + "…"


def pdf(src, out_dir, head):
    import pymupdf as fitz

    doc = fitz.open(src)
    txt = out_path(src, out_dir, ".txt")
    low, with_img, chars = [], [], 0
    with open(txt, "w", encoding="utf-8") as f:
        for i, page in enumerate(doc, 1):
            t = page.get_text()
            chars += len(t)
            f.write(f"\n=== PAGE {i} ===\n{t}")
            if len(t.strip()) < 50:
                low.append(i)
            if page.get_images():
                with_img.append(i)
    print(f"pages: {doc.page_count}  chars: {chars}  title: {clip(doc.metadata.get('title') or '-')}")
    if low:
        print(f"pages with <50 chars of text (scanned or image-only): {fmt(low)}")
    if with_img:
        print(f"pages with embedded images (check chart/table if text is thin): {fmt(with_img)}")
    if len(low) > doc.page_count / 2:
        print("mostly scanned: text is unreliable, OCR needed (tesseract, if installed) or visual check by subagent")
    print(f"full text: {txt}")
    first = [l for l in Path(txt).read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("=== PAGE")]
    show(first, head)


def fmt(nums, cap=20):
    s = ",".join(map(str, nums[:cap]))
    return s + (f" (+{len(nums) - cap} more)" if len(nums) > cap else "")


def show(lines, head):
    for l in lines[:head]:
        print("  | " + clip(l))
    if len(lines) > head:
        print(f"  ({len(lines) - head} more lines not shown)")


def xlsx(src, out_dir, head):
    import openpyxl

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    print(f"sheets: {len(wb.sheetnames)}")
    csv_dir = out_path(src, out_dir, "")
    csv_dir.mkdir(exist_ok=True)
    for ws in wb.worksheets:
        rows = 0
        sample = []
        target = csv_dir / (ws.title.replace("/", "_")[:40] + ".csv")
        with open(target, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                rows += 1
                w.writerow(["" if c is None else c for c in row])
                if len(sample) < head:
                    sample.append(row)
        cols = max((len(r) for r in sample), default=0)
        print(f"- {ws.title!r}: {rows} rows x {cols} cols  csv: {target}")
        for r in sample:
            print("  | " + clip(" ; ".join("" if c is None else str(c) for c in r)))
    print("note: values are cached results (data_only); formulas never evaluated")


def csvfile(src, out_dir, head):
    with open(src, encoding="utf-8", errors="replace") as f:
        lines = f.read().splitlines()
    print(f"lines: {len(lines)}  bytes: {src.stat().st_size}")
    show(lines, head)


def docx(src, out_dir, head):
    import docx as D

    d = D.Document(src)
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    txt = out_path(src, out_dir, ".txt")
    with open(txt, "w", encoding="utf-8") as f:
        f.write("\n".join(paras))
        for ti, t in enumerate(d.tables, 1):
            f.write(f"\n=== TABLE {ti} ===\n")
            for r in t.rows:
                f.write(" | ".join(c.text.strip() for c in r.cells) + "\n")
    print(f"paragraphs: {len(paras)}  tables: {len(d.tables)}")
    print(f"full text: {txt}")
    show(paras, head)


def archive(src, out_dir, head):
    ext = src.suffix.lower()
    if ext == ".rar":
        tool = shutil.which("unrar")
        cmd = [tool, "lb", "--", str(src)] if tool else None
        note = "" if tool else "unrar missing (7z 16 and unrar-free cannot extract RAR5): run install.sh"
    else:
        tool = shutil.which("7z")
        cmd = [tool, "l", "-ba", "--", str(src)] if tool else None
        note = "" if tool else "7z missing: run install.sh"
    if not cmd:
        print(note)
        return
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
    lines = [l for l in r.stdout.splitlines() if l.strip()]
    print(f"size: {src.stat().st_size} bytes  entries listed: {len(lines)}")
    if r.returncode:
        print(f"listing failed rc={r.returncode}: {clip(r.stderr)}")
    show(lines, head)
    print("not extracted. Extract into a NEW EMPTY dir: unrar x -idq -y FILE DIR/  |  7z x -y -bso0 -oDIR FILE")


def image(src, out_dir, head):
    try:
        import pymupdf as fitz

        p = fitz.Pixmap(str(src))
        print(f"image: {p.width}x{p.height}  {src.stat().st_size} bytes")
    except Exception as e:
        print(f"image: {src.stat().st_size} bytes ({e})")
    print("no text layer: look at it only if needed, preferably through a subagent that returns the conclusion")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--out", default="/tmp/ler-documento")
    ap.add_argument("--head", type=int, default=15)
    a = ap.parse_args()
    src = Path(a.file)
    if not src.is_file():
        sys.exit(f"not a file: {src}")
    out = Path(a.out)
    ext = src.suffix.lower()
    print(f"file: {src.name}  ({src.stat().st_size} bytes)")
    if ext == ".pdf":
        try:
            pdf(src, out, a.head)
        except Exception as e:
            sys.exit(f"cannot open as PDF ({clip(e)}); check `file` and `head -c 200`")
    elif ext in {".xlsx", ".xlsm"}:
        xlsx(src, out, a.head)
    elif ext in {".csv", ".tsv", ".txt", ".md", ".log", ".json"}:
        csvfile(src, out, a.head)
    elif ext == ".docx":
        docx(src, out, a.head)
    elif ext in ARCH:
        archive(src, out, a.head)
    elif ext in IMG:
        image(src, out, a.head)
    elif ext in MEDIA:
        print("audio/video: use the transcrever-audio skill")
    else:
        print(f"unknown type {ext!r}; try `file` and `head -c 300`")


if __name__ == "__main__":
    main()
